import os
import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from config import DEFAULT_EXCEL_PATH

class ExcelLogger:
    def __init__(self, excel_path: str = DEFAULT_EXCEL_PATH):
        self.excel_path = excel_path
        self.headers = [
            "Số thứ tự",
            "Tên file",
            "Chủ đề",
            "Nội dung thuế liên quan",
            "Link gốc",
            "Thời lượng cắt"
        ]
        self._ensure_file_exists()

    def _ensure_file_exists(self):
        """
        Khởi tạo file Excel mới nếu chưa tồn tại, kèm định dạng đẹp mắt.
        """
        output_dir = os.path.dirname(self.excel_path)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir, exist_ok=True)

        if not os.path.exists(self.excel_path):
            wb = Workbook()
            ws = wb.active
            ws.title = "Danh Sách Video Thuế"

            # Style cho Tiêu đề (Header)
            header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

            ws.append(self.headers)
            for col_num in range(1, len(self.headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center

            # Đặt độ rộng chuẩn cho các cột
            col_widths = {
                "A": 10,  # Số thứ tự
                "B": 35,  # Tên file
                "C": 30,  # Chủ đề
                "D": 65,  # Nội dung thuế liên quan
                "E": 45,  # Link gốc
                "F": 18   # Thời lượng cắt
            }
            for col, width in col_widths.items():
                ws.column_dimensions[col].width = width

            wb.save(self.excel_path)
            wb.close()
            print(f"[EXCEL] Khởi tạo file Excel mới tại: {self.excel_path}")

    def append_log(
        self,
        filename: str,
        topic: str,
        tax_content: str,
        original_link: str,
        duration_str: str
    ):
        """
        Ghi thêm 1 dòng dữ liệu mới vào cuối file Excel.
        """
        wb = load_workbook(self.excel_path)
        ws = wb.active

        # Xử lý STT tự động tăng
        max_row = ws.max_row
        stt = max_row if max_row > 1 else 1

        new_row = [
            stt,
            filename,
            topic,
            tax_content,
            original_link,
            duration_str
        ]

        ws.append(new_row)
        current_row = ws.max_row

        # Format dòng mới
        font_regular = Font(name="Segoe UI", size=10)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_center = Alignment(horizontal="center", vertical="center")

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for col_idx in range(1, len(new_row) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = font_regular
            cell.border = thin_border

            if col_idx in [1, 6]: # STT, Thời lượng
                cell.alignment = align_center
            else:
                cell.alignment = align_left

        wb.save(self.excel_path)
        wb.close()
        print(f"[EXCEL LOGGED] Đã thêm thành công dòng {stt} cho file {filename}")

